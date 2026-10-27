"""Step 4: a data-quality digest for each partner.

In the real project these went to partners every day. Each workbook holds what
a partner's field coordinator needs: a summary, the flagged interviews to
follow up, the interviews removed and why, an enumerator table, and progress
against target in each of its LGAs. Every sheet carries the synthetic-data label.

Inputs: data/synthetic/raw/, data/synthetic/clean/
Outputs: data/synthetic/partner_digests/<partner>.xlsx
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "python"))
from monitoring_tools import SYNTHETIC_LABEL, load_config, make_xlsx_reproducible, output_dir  # noqa: E402

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LABEL_FONT = Font(bold=True, color="C00000", size=12)


def write_sheet(writer, name, df, title):
    df.to_excel(writer, sheet_name=name, index=False, startrow=3)
    ws = writer.sheets[name]
    ws["A1"], ws["A2"] = title, SYNTHETIC_LABEL
    ws["A1"].font, ws["A2"].font = Font(bold=True, size=14), LABEL_FONT
    for cell in ws[4]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    ws.freeze_panes = "A5"
    for i, col in enumerate(df.columns, start=1):
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
        ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)


def main() -> int:
    cfg = load_config()
    out = output_dir(cfg)
    raw = pd.concat([pd.read_csv(f, dtype=str) for f in sorted((out / "raw").glob("submissions_*.csv.gz"))])
    clean = pd.concat([pd.read_csv(f, dtype=str) for f in sorted((out / "clean").glob("clean_data_*.csv.gz"))])
    log = pd.read_csv(out / "clean" / "cleaning_log.csv", dtype=str)
    deleted = pd.read_csv(out / "clean" / "deletion_log.csv", dtype=str)
    enum = pd.read_csv(out / "clean" / "enumerator_flags.csv")
    achieved = pd.read_csv(out / "clean" / "achieved_by_stratum.csv")
    as_of = raw["today"].max()

    dest = out / "partner_digests"
    dest.mkdir(exist_ok=True)
    for old in dest.glob("*.xlsx"):
        old.unlink()

    for code in sorted(raw["partner"].unique()):
        name = code.replace("_", " ").title()
        r, c = raw[raw["partner"] == code], clean[clean["partner"] == code]
        lg, dl = log[log["partner"] == code], deleted[deleted["partner"] == code]
        flags = lg[lg["action"] == "flag"]
        ach = achieved[achieved["partner"] == name]
        summary = pd.DataFrame({
            "Item": ["Partner", "States", "Data as of", "Submissions received", "Interviews kept",
                     "Interviews removed", "Values corrected", "Interviews flagged for follow-up",
                     "Households targeted", "Households achieved", "Progress against target"],
            "Value": [name, ", ".join(sorted(c["adm1_name"].dropna().unique())), as_of, len(r), len(c), len(dl),
                      int((lg["action"] == "change").sum()), flags["_uuid"].nunique(), int(ach["target"].sum()),
                      int(ach["achieved"].sum()),
                      f"{100 * ach['achieved'].sum() / max(1, ach['target'].sum()):.1f}%"]})
        follow_up = flags.merge(raw[["_uuid", "today", "cluster_id", "household_id"]].drop_duplicates("_uuid"),
                                on="_uuid", how="left")[["today", "enumerator_id", "cluster_id", "household_id",
                                                         "_uuid", "check", "issue", "old_value"]] \
            .rename(columns={"today": "date", "old_value": "value"}).sort_values(["date", "enumerator_id"])
        removed = dl.merge(raw[["_index", "today", "cluster_id", "household_id"]], on="_index", how="left")[
            ["today", "enumerator_id", "cluster_id", "household_id", "_uuid", "check_id", "reason"]] \
            .rename(columns={"today": "date"}).sort_values(["date", "enumerator_id"])
        e = enum[enum["partner"] == code].copy()
        e["removed"] = e["enumerator_id"].map(dl["enumerator_id"].value_counts()).fillna(0).astype(int)
        e["flagged"] = e["enumerator_id"].map(flags.drop_duplicates("_uuid")["enumerator_id"].value_counts()) \
            .fillna(0).astype(int)
        e["pattern"] = [", ".join(p for p, on in [("much faster than others", fast), ("repeats the same answers", same)]
                                  if on) for fast, same in zip(e["flag_fast"], e["flag_identical"])]
        enumerators = e.assign(median_duration=e["median_duration"].round(1),
                               identical_share=(100 * e["identical_share"]).round(1))[
            ["enumerator_id", "interviews", "median_duration", "identical_share", "removed", "flagged", "pattern"]] \
            .rename(columns={"median_duration": "median minutes", "identical_share": "% repeated answers",
                             "removed": "interviews removed", "flagged": "interviews flagged"})
        progress = ach[["adm1_name", "adm2_name", "pop_type", "target", "achieved", "pct"]] \
            .rename(columns={"pct": "percent"}).sort_values(["adm1_name", "adm2_name", "pop_type"])

        path = dest / f"{code}.xlsx"
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            write_sheet(writer, "Summary", summary, f"Data-quality digest: {name}")
            write_sheet(writer, "Follow up", follow_up, "Interviews flagged: please check with the enumerator")
            write_sheet(writer, "Removed", removed, "Interviews removed from the dataset, and why")
            write_sheet(writer, "Enumerators", enumerators, "Enumerators")
            write_sheet(writer, "Progress", progress, "Progress against target by LGA")
        make_xlsx_reproducible(path)

    print(f"Partner digests: {raw['partner'].nunique()} written to {cfg['outputs']['dir']}/partner_digests (as of {as_of})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
